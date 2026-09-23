# -*- coding: utf-8 -*-
"""
Genera las "Listas Descargables": un Excel por Zona con el detalle de stock
disponible de cada tienda de esa zona en una hoja separada.

Reemplaza el proceso manual (armar el Excel a mano para cada supervisor que
lo pide) por uno automatico, en el mismo formato usado la primera vez
("Obsolescencia Zona RM2.xlsx" con una hoja por tienda).

Reutiliza el mismo procesador de datos que el portal y el dashboard interno
(dashboard-obsolescencia/src/services/cargador_excel.py), asi que las cifras
son siempre las mismas que ve el publico en el portal de consultas.

Reglas:
- Solo productos con stock > 0 (se excluye lo agotado).
- Una hoja por tienda de la zona, nombrada solo con el nombre de la tienda
  (sin el codigo AP0000/SG0000).
- Un archivo .xlsx por zona: "Obsolescencia <Zona>.xlsx".
- El archivo se sobreescribe cada vez que se corre el script (siempre
  refleja el stock actual del maestro).

Cada archivo se guarda en DOS lugares:
  - "Listas Descargables/" (uso manual, para compartir directo)
  - "portal-web/descargas/" (se publica junto con index.html: el boton
    "Descargar zona completa" del portal apunta ahi)

generar_portal.py llama a generar(resultado) reutilizando el Excel que ya
leyo, para no releer el archivo maestro dos veces.

Ejecutar desde esta carpeta:
    python generar_listas_descargables.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

RAIZ_PORTAL = Path(__file__).resolve().parent
RAIZ_PROYECTO = RAIZ_PORTAL.parent
RAIZ_STREAMLIT = RAIZ_PROYECTO / "dashboard-obsolescencia"
sys.path.insert(0, str(RAIZ_STREAMLIT))

from src.services import cargador_excel as datos  # noqa: E402
from src.utils.formato import dividir_tienda, etiqueta_tienda  # noqa: E402

CARPETA_SALIDA = RAIZ_PROYECTO / "Listas Descargables"
CARPETA_SALIDA_PORTAL = RAIZ_PORTAL / "descargas"

# (campo interno del df, encabezado de la columna en el Excel de salida)
COLUMNAS = [
    ("zona", "Zona"),
    ("_tienda_label", "Tienda"),
    ("material", "Material"),
    ("codigo_fabricante", "Código fabricante"),
    ("texto_breve", "Texto breve"),
    ("marca", "Marca"),
    ("categoria", "Categoría"),
    ("subcategoria", "Subcategoría"),
    ("marca_vehiculo", "Marca vehículo"),
    ("app_modelo", "Modelo"),
    ("app_motor", "Motor"),
    ("app_anios", "Rango de años original"),
    ("stock", "Stock"),
    ("precio_normal", "Precio normal"),
    ("valor_remate", "Precio Liquidación"),
]
CAMPOS_NUMERICOS = {"material", "stock", "precio_normal", "valor_remate"}

_CARACTERES_INVALIDOS_HOJA = re.compile(r'[\\/*?:\[\]]')
_CARACTERES_INVALIDOS_ARCHIVO = re.compile(r'[\\/*?:"<>|]')


def _valor_celda(valor, campo: str):
    """Convierte NaN/None a celda vacia y deja numeros como numeros (sin
    decimales de sobra), igual que en el Excel original a mano."""
    if valor is None:
        return None
    try:
        if valor != valor:  # NaN
            return None
    except TypeError:
        pass
    if campo in CAMPOS_NUMERICOS:
        try:
            numero = float(valor)
            return int(numero) if numero == int(numero) else numero
        except (TypeError, ValueError):
            return valor
    return valor


def _nombre_hoja(nombre: str, usados: set[str]) -> str:
    limpio = _CARACTERES_INVALIDOS_HOJA.sub(" ", nombre).strip() or "Tienda"
    limpio = limpio[:31]
    base, i = limpio, 2
    while limpio.lower() in usados:
        sufijo = f" ({i})"
        limpio = base[: 31 - len(sufijo)] + sufijo
        i += 1
    usados.add(limpio.lower())
    return limpio


_ANCHO_MAXIMO_COLUMNA = 45


def _formatear_hoja(ws) -> None:
    """Encabezado en negrita, fila fija y ancho de columna segun contenido
    (solo estetica: no cambia ningun dato)."""
    for celda in ws[1]:
        celda.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, columna in enumerate(ws.columns, start=1):
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
        ws.column_dimensions[get_column_letter(i)].width = min(ancho + 2, _ANCHO_MAXIMO_COLUMNA)


def _nombre_archivo(zona: str) -> str:
    # Mismo criterio usado a mano la primera vez: "Zona RM 2" -> "RM2" en el
    # nombre del archivo (el valor de la columna Zona adentro no cambia).
    compacto = re.sub(r"(RM)\s+(\d)", r"\1\2", zona)
    limpio = _CARACTERES_INVALIDOS_ARCHIVO.sub("-", compacto).strip()
    return f"Obsolescencia {limpio}.xlsx"


def generar(resultado=None):
    resultado = resultado or datos.cargar()
    df = resultado.df

    disponible = df[df["stock"] > 0].copy()
    excluidos_sin_stock = len(df) - len(disponible)

    partes = disponible["tienda"].map(dividir_tienda)
    disponible["_tienda_codigo"] = [c for c, n in partes]
    disponible["_tienda_nombre"] = [n or c for c, n in partes]
    disponible["_tienda_label"] = [etiqueta_tienda(c, n) for c, n in partes]

    CARPETA_SALIDA.mkdir(exist_ok=True)
    CARPETA_SALIDA_PORTAL.mkdir(exist_ok=True)

    resumen = []
    for zona in sorted(disponible["zona"].unique()):
        df_zona = disponible[disponible["zona"] == zona]

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        usados: set[str] = set()

        tiendas = (
            df_zona[["_tienda_codigo", "_tienda_nombre"]]
            .drop_duplicates()
            .sort_values("_tienda_codigo")
            .values.tolist()
        )

        filas_zona = 0
        for codigo, nombre in tiendas:
            df_tienda = df_zona[df_zona["_tienda_codigo"] == codigo]
            ws = wb.create_sheet(_nombre_hoja(nombre, usados))
            ws.append([etiqueta for _, etiqueta in COLUMNAS])
            for _, fila in df_tienda.iterrows():
                ws.append([
                    _valor_celda(fila.get(campo), campo) for campo, _ in COLUMNAS
                ])
            _formatear_hoja(ws)
            filas_zona += len(df_tienda)

        nombre_archivo = _nombre_archivo(zona)
        wb.save(CARPETA_SALIDA / nombre_archivo)
        wb.save(CARPETA_SALIDA_PORTAL / nombre_archivo)
        resumen.append((zona, len(tiendas), filas_zona, nombre_archivo))

    return resultado, resumen, excluidos_sin_stock


def imprimir_reporte(resultado, resumen, excluidos_sin_stock) -> None:
    print(f"Base usada: {resultado.archivo.name}")
    print(f"Productos excluidos por stock 0 (o sin stock legible): {excluidos_sin_stock}")
    print(f"\nArchivos generados en «{CARPETA_SALIDA}» y en «{CARPETA_SALIDA_PORTAL}»:")
    total_filas = 0
    for zona, n_tiendas, filas, nombre_archivo in resumen:
        print(f"  - {nombre_archivo}: {n_tiendas} tienda(s), {filas} filas")
        total_filas += filas
    print(f"\nTotal: {len(resumen)} archivos, {total_filas} filas de stock disponible.")
    if resultado.avisos:
        print(f"\nAvisos de la base ({len(resultado.avisos)}):")
        for a in resultado.avisos:
            print(f"  - {a}")


if __name__ == "__main__":
    resultado, resumen, excluidos_sin_stock = generar()
    imprimir_reporte(resultado, resumen, excluidos_sin_stock)
